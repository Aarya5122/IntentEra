"""Reusable diagram primitives: icon-decorated nodes, arrows, lane bands."""
from __future__ import annotations

from dataclasses import dataclass
from typing import Optional

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from theme import (
    CYAN, INDIGO, VIOLET, TEAL, GREEN, AMBER, RED,
    PANEL, PANEL_EDGE, BG_DEEP, BG_MID,
    TEXT, TEXT_DIM, TEXT_MUTED, WHITE,
    FONT_HEAD, FONT_BODY, FONT_MONO,
    add_text, fill_gradient, asset as asset_path,
)


# ----------------------------------------------------------------------
# Geometry of a node
# ----------------------------------------------------------------------

NODE_W = Inches(1.8)
NODE_H = Inches(1.35)
ICON_SIZE = Inches(0.55)


@dataclass
class Node:
    x: Emu
    y: Emu
    w: Emu
    h: Emu
    shape: object        # the rounded-rectangle shape
    anchor_top: tuple[int, int]
    anchor_bottom: tuple[int, int]
    anchor_left: tuple[int, int]
    anchor_right: tuple[int, int]
    center: tuple[int, int]


def _anchors(x, y, w, h) -> dict:
    cx = x + w // 2
    cy = y + h // 2
    return {
        "anchor_top": (cx, y),
        "anchor_bottom": (cx, y + h),
        "anchor_left": (x, cy),
        "anchor_right": (x + w, cy),
        "center": (cx, cy),
    }


def node(
    slide,
    x,
    y,
    *,
    label: str,
    sublabel: str | None = None,
    icon: str | None = None,      # e.g. "aws/Lambda" or "logos/mongodb"
    w=NODE_W,
    h=NODE_H,
    fill: RGBColor = PANEL,
    border: RGBColor = PANEL_EDGE,
    accent: RGBColor = CYAN,
    label_size: int = 12,
    sublabel_size: int = 9,
    corner: float = 0.12,
) -> Node:
    # Card
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try:
        card.adjustments[0] = corner
    except Exception:
        pass
    card.fill.solid()
    card.fill.fore_color.rgb = fill
    card.line.color.rgb = border
    card.line.width = Pt(0.75)
    card.shadow.inherit = False

    # Top accent strip
    strip = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, x, y, w, Emu(38100)  # ~4px
    )
    strip.line.fill.background()
    fill_gradient(strip, [(0, accent), (1, INDIGO)], angle=0)

    # Icon
    if icon:
        kind, name = icon.split("/", 1)
        path = asset_path(kind, name)
        icon_w = ICON_SIZE
        icon_h = ICON_SIZE
        ix = x + (w - icon_w) // 2
        iy = y + Inches(0.18)
        slide.shapes.add_picture(path, ix, iy, icon_w, icon_h)
        text_y = iy + icon_h + Inches(0.06)
    else:
        text_y = y + Inches(0.25)

    # Label
    add_text(
        slide, x + Inches(0.1), text_y, w - Inches(0.2), Inches(0.3),
        label, size=label_size, color=TEXT, font=FONT_HEAD,
        bold=True, align=PP_ALIGN.CENTER,
    )

    # Sublabel
    if sublabel:
        add_text(
            slide, x + Inches(0.1), text_y + Inches(0.3),
            w - Inches(0.2), Inches(0.28),
            sublabel, size=sublabel_size, color=TEXT_MUTED,
            font=FONT_MONO, align=PP_ALIGN.CENTER,
        )

    return Node(x, y, w, h, card, **_anchors(x, y, w, h))


# ----------------------------------------------------------------------
# Arrows / connectors
# ----------------------------------------------------------------------

def arrow(
    slide,
    frm: Node | tuple[int, int],
    to: Node | tuple[int, int],
    *,
    side_from: str = "right",
    side_to: str = "left",
    color: RGBColor = CYAN,
    dashed: bool = False,
    width: float = 1.75,
    head: bool = True,
    label: str | None = None,
    label_above: bool = True,
):
    if isinstance(frm, Node):
        fx, fy = getattr(frm, f"anchor_{side_from}")
    else:
        fx, fy = frm
    if isinstance(to, Node):
        tx, ty = getattr(to, f"anchor_{side_to}")
    else:
        tx, ty = to

    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, fx, fy, tx, ty)
    conn.line.color.rgb = color
    conn.line.width = Pt(width)

    ln = conn.line._get_or_add_ln()
    from lxml import etree
    from theme import _qn

    # dash style
    if dashed:
        existing = ln.find(_qn("a:prstDash"))
        if existing is not None:
            ln.remove(existing)
        etree.SubElement(ln, _qn("a:prstDash"), {"val": "dash"})

    # arrow head
    if head:
        for tag in ("a:headEnd", "a:tailEnd"):
            el = ln.find(_qn(tag))
            if el is not None:
                ln.remove(el)
        etree.SubElement(ln, _qn("a:tailEnd"),
                         {"type": "triangle", "w": "med", "len": "med"})

    # Optional label
    if label:
        mx = (fx + tx) // 2
        my = (fy + ty) // 2
        offset = Inches(-0.22) if label_above else Inches(0.06)
        add_text(
            slide, mx - Inches(0.9), my + offset, Inches(1.8), Inches(0.25),
            label, size=9, color=TEXT_DIM, font=FONT_MONO,
            align=PP_ALIGN.CENTER,
        )
    return conn


def elbow(
    slide,
    frm: Node | tuple[int, int],
    to: Node | tuple[int, int],
    *,
    side_from: str = "bottom",
    side_to: str = "top",
    color: RGBColor = CYAN,
    dashed: bool = False,
    width: float = 1.75,
    head: bool = True,
):
    if isinstance(frm, Node):
        fx, fy = getattr(frm, f"anchor_{side_from}")
    else:
        fx, fy = frm
    if isinstance(to, Node):
        tx, ty = getattr(to, f"anchor_{side_to}")
    else:
        tx, ty = to

    conn = slide.shapes.add_connector(MSO_CONNECTOR.ELBOW, fx, fy, tx, ty)
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    ln = conn.line._get_or_add_ln()
    from lxml import etree
    from theme import _qn
    if dashed:
        etree.SubElement(ln, _qn("a:prstDash"), {"val": "dash"})
    if head:
        etree.SubElement(ln, _qn("a:tailEnd"),
                         {"type": "triangle", "w": "med", "len": "med"})
    return conn


# ----------------------------------------------------------------------
# Lane / container
# ----------------------------------------------------------------------

def lane(
    slide,
    x,
    y,
    w,
    h,
    *,
    title: str,
    subtitle: str | None = None,
    color: RGBColor = INDIGO,
    alpha_fill: RGBColor | None = None,
):
    """A titled container with a colored header stripe."""
    bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    bg.adjustments[0] = 0.04
    bg.fill.solid()
    bg.fill.fore_color.rgb = alpha_fill or BG_MID
    bg.line.color.rgb = color
    bg.line.width = Pt(1.0)

    # Header pill
    pill_w = Inches(2.4)
    pill = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.2), y - Inches(0.18),
        pill_w, Inches(0.36),
    )
    pill.adjustments[0] = 0.5
    pill.fill.solid()
    pill.fill.fore_color.rgb = color
    pill.line.fill.background()
    tf = pill.text_frame
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.name = FONT_HEAD
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = BG_DEEP

    if subtitle:
        add_text(
            slide, x + Inches(0.2) + pill_w + Inches(0.15),
            y - Inches(0.14), Inches(4), Inches(0.3),
            subtitle, size=10, color=TEXT_MUTED, font=FONT_MONO,
        )
    return bg


# ----------------------------------------------------------------------
# Layered architecture band (for slide 9)
# ----------------------------------------------------------------------

def layer_band(
    slide, x, y, w, h, *,
    title: str, items: list[str],
    color: RGBColor = INDIGO, icon: str | None = None,
):
    bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    bg.adjustments[0] = 0.12
    bg.fill.solid()
    bg.fill.fore_color.rgb = PANEL
    bg.line.color.rgb = color
    bg.line.width = Pt(1.0)

    # Left color stripe
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.12), h)
    stripe.line.fill.background()
    fill_gradient(stripe, [(0, color), (1, CYAN)], angle=5400000)

    inner_x = x + Inches(0.35)
    # Title
    add_text(slide, inner_x, y + Inches(0.08), Inches(3.2), Inches(0.35),
             title, size=15, color=TEXT, font=FONT_HEAD, bold=True)

    # Items on the right side
    items_x = x + Inches(3.7)
    add_text(slide, items_x, y + Inches(0.08), w - Inches(3.9), h - Inches(0.2),
             "    ·    ".join(items),
             size=11, color=TEXT_DIM, font=FONT_BODY,
             line_spacing=1.25, anchor=MSO_ANCHOR.MIDDLE)
    return bg
