"""Visual theme: colors, fonts, gradient backgrounds, and slide chrome.

All colors in sRGB hex. Geometry uses python-pptx `Emu` / `Inches` units.
Slides are 16:9 at 13.333in x 7.5in (1280 x 720 at 96dpi).
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from lxml import etree

ROOT = Path(__file__).parent
ASSETS = ROOT / "assets"

# --- Slide geometry ---
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# --- Palette (tech gradient: indigo -> violet -> cyan) ---
BG_DEEP = RGBColor(0x0B, 0x10, 0x26)       # deep indigo
BG_MID = RGBColor(0x11, 0x1B, 0x3F)        # indigo mid
BG_SOFT = RGBColor(0x0F, 0x2A, 0x4A)       # blueish accent end
PANEL = RGBColor(0x17, 0x23, 0x4A)         # card body
PANEL_EDGE = RGBColor(0x2E, 0x3F, 0x77)    # card border
INDIGO = RGBColor(0x63, 0x66, 0xF1)
VIOLET = RGBColor(0xA7, 0x8B, 0xFA)
CYAN = RGBColor(0x22, 0xD3, 0xEE)
TEAL = RGBColor(0x2D, 0xD4, 0xBF)
AMBER = RGBColor(0xFB, 0xBF, 0x24)
GREEN = RGBColor(0x34, 0xD3, 0x99)
RED = RGBColor(0xF8, 0x71, 0x71)
TEXT = RGBColor(0xF8, 0xFA, 0xFC)
TEXT_DIM = RGBColor(0xCB, 0xD5, 0xE1)
TEXT_MUTED = RGBColor(0x94, 0xA3, 0xB8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# --- Fonts ---
FONT_HEAD = "Montserrat"       # falls back on Calibri / system sans
FONT_BODY = "Inter"
FONT_MONO = "JetBrains Mono"


# ----------------------------------------------------------------------
# Low-level XML helpers (python-pptx doesn't expose gradient fills nicely)
# ----------------------------------------------------------------------

_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
NSMAP = {"a": _A_NS}


def _qn(tag: str) -> str:
    prefix, local = tag.split(":")
    return f"{{{_A_NS}}}{local}"


def _hex(c: RGBColor) -> str:
    return "{:02X}{:02X}{:02X}".format(c[0], c[1], c[2])


def _get_spPr(shape):
    # shape._element is <p:sp>, which contains <p:spPr>
    return shape._element.spPr


def fill_gradient(shape, stops, angle: int = 2700000) -> None:
    """Apply a linear gradient to any shape that has a spPr element.

    stops: list of (position_0_1, RGBColor).
    angle: 1/60000 degree units (2700000 = 45deg).
    """
    spPr = _get_spPr(shape)
    # Remove existing fill children
    for tag in ("a:solidFill", "a:gradFill", "a:noFill", "a:pattFill", "a:blipFill"):
        el = spPr.find(_qn(tag))
        if el is not None:
            spPr.remove(el)

    # Must appear before <a:ln>, <a:effectLst> inside <p:spPr>.
    grad = etree.Element(_qn("a:gradFill"),
                         {"flip": "none", "rotWithShape": "1"})
    gsLst = etree.SubElement(grad, _qn("a:gsLst"))
    for pos, color in stops:
        gs = etree.SubElement(gsLst, _qn("a:gs"),
                              {"pos": str(int(pos * 100000))})
        etree.SubElement(gs, _qn("a:srgbClr"), {"val": _hex(color)})
    etree.SubElement(grad, _qn("a:lin"),
                     {"ang": str(angle), "scaled": "0"})
    etree.SubElement(grad, _qn("a:tileRect"))

    # Find insertion index: before <a:ln> if present, else append.
    ln = spPr.find(_qn("a:ln"))
    if ln is not None:
        spPr.insert(list(spPr).index(ln), grad)
    else:
        spPr.append(grad)
    return shape


# ----------------------------------------------------------------------
# Slide chrome builders
# ----------------------------------------------------------------------

def add_background(slide, variant: str = "content") -> None:
    """Paint a full-bleed gradient background rectangle at z-index 0."""
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    rect.line.fill.background()
    if variant == "title":
        stops = [(0, BG_DEEP), (0.55, RGBColor(0x18, 0x17, 0x64)), (1, RGBColor(0x06, 0x3F, 0x4E))]
    elif variant == "section":
        stops = [(0, BG_DEEP), (1, RGBColor(0x30, 0x27, 0x80))]
    else:  # content
        stops = [(0, BG_DEEP), (1, BG_MID)]
    fill_gradient(rect, stops, angle=2700000)
    # send to back: move to index 0
    spTree = rect._element.getparent()
    spTree.remove(rect._element)
    spTree.insert(2, rect._element)


def add_accent_bar(slide) -> None:
    """Thin cyan accent bar on the left edge."""
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.08), SLIDE_H)
    bar.line.fill.background()
    fill_gradient(bar, [(0, CYAN), (1, INDIGO)], angle=5400000)  # vertical


def add_footer(slide, page_no: int, total: int, section: str = "") -> None:
    # bottom hairline
    hair = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                  Inches(0.6), Inches(7.10),
                                  Inches(12.1), Emu(9525))  # ~1px
    hair.line.fill.background()
    hair.fill.solid()
    hair.fill.fore_color.rgb = PANEL_EDGE

    # left meta
    add_text(slide, Inches(0.6), Inches(7.15), Inches(8), Inches(0.3),
             "IntentEra  ·  Midterm Review  ·  Aarya Nanndaann Singh M N (2024MT03013)",
             size=10, color=TEXT_MUTED, font=FONT_BODY)

    # right page
    add_text(slide, Inches(11.8), Inches(7.15), Inches(1.5), Inches(0.3),
             f"{page_no:02d} / {total:02d}",
             size=10, color=TEXT_MUTED, font=FONT_MONO, align=PP_ALIGN.RIGHT)


def add_title(slide, text: str, eyebrow: str | None = None) -> None:
    """Big page title with optional small eyebrow label above it."""
    y = Inches(0.45)
    if eyebrow:
        add_text(slide, Inches(0.6), y, Inches(12), Inches(0.3),
                 eyebrow.upper(),
                 size=11, color=CYAN, font=FONT_HEAD, bold=True, spacing=200)
        y = Inches(0.78)
    add_text(slide, Inches(0.6), y, Inches(12.2), Inches(0.85),
             text,
             size=30, color=TEXT, font=FONT_HEAD, bold=True)

    # gradient underline
    underline = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                       Inches(0.6), Inches(1.6),
                                       Inches(2.2), Emu(38100))  # ~4px
    underline.line.fill.background()
    fill_gradient(underline, [(0, CYAN), (1, INDIGO)], angle=0)


# ----------------------------------------------------------------------
# Text & panel utilities
# ----------------------------------------------------------------------

def add_text(slide, x, y, w, h, text: str, *,
             size: int = 14, color: RGBColor = TEXT,
             font: str = FONT_BODY, bold: bool = False,
             italic: bool = False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, spacing: int | None = None,
             line_spacing: float | None = None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = Inches(0)
    tf.margin_top = tf.margin_bottom = Inches(0)
    tf.word_wrap = True
    tf.vertical_anchor = anchor

    # first paragraph already exists
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing is not None:
            p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        if spacing is not None:
            # character spacing in 1/100th point
            rPr = run._r.get_or_add_rPr()
            rPr.set("spc", str(spacing))
    return tb


def add_panel(slide, x, y, w, h, *,
              fill: RGBColor = PANEL, border: RGBColor = PANEL_EDGE,
              corner: float = 0.15, glow: bool = False):
    """Rounded-rectangle card."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    # adjust corner radius
    try:
        shape.adjustments[0] = corner
    except Exception:
        pass
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    return shape


def add_chip(slide, x, y, label: str, *, color: RGBColor = CYAN,
             text_color: RGBColor | None = None,
             width: Inches | None = None):
    """Small pill-shaped label."""
    w = width or Inches(1.2)
    h = Inches(0.32)
    chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    chip.adjustments[0] = 0.5
    chip.fill.solid()
    chip.fill.fore_color.rgb = color
    chip.line.fill.background()
    tf = chip.text_frame
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = FONT_HEAD
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = text_color or BG_DEEP
    return chip


def add_bullet(slide, x, y, w, h, items, *,
               size: int = 14, color: RGBColor = TEXT_DIM,
               bullet: str = "•", bullet_color: RGBColor | None = None,
               line_spacing: float = 1.3):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = Inches(0)
    tf.margin_top = tf.margin_bottom = Inches(0)
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = line_spacing
        b = p.add_run()
        b.text = f"{bullet}  "
        b.font.name = FONT_BODY
        b.font.size = Pt(size)
        b.font.bold = True
        b.font.color.rgb = bullet_color or CYAN

        r = p.add_run()
        r.text = item
        r.font.name = FONT_BODY
        r.font.size = Pt(size)
        r.font.color.rgb = color
    return tb


def asset(kind: str, name: str) -> str:
    """Return str path to an asset (for add_picture)."""
    p = ASSETS / kind / f"{name}.png"
    if not p.exists():
        raise FileNotFoundError(p)
    return str(p)
