"""Lightweight native-PPTX diagram primitives.

The major diagrams are pre-rendered PNGs under Final Report/figures/.
This module is only used for small inline diagrams we draw on the slide:
- the formal-problem callout (R, C, D, q -> E, S)
- the multi-query fan-out mini-diagram on the retriever slide
- the layered architecture band (decorative)
"""
from __future__ import annotations

from dataclasses import dataclass

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from theme import (
    AMBER, BG_DEEP, BG_MID, CYAN, FONT_BODY, FONT_HEAD, FONT_MONO,
    GREEN, INDIGO, PANEL, PANEL_EDGE, RED, TEAL, TEXT, TEXT_DIM,
    TEXT_MUTED, VIOLET, WHITE,
    _qn, add_text, fill_gradient,
)


# ----------------------------------------------------------------------
# Node geometry
# ----------------------------------------------------------------------

NODE_W = Inches(1.85)
NODE_H = Inches(0.95)


@dataclass
class Node:
    x: Emu
    y: Emu
    w: Emu
    h: Emu
    shape: object
    anchor_top: tuple[int, int]
    anchor_bottom: tuple[int, int]
    anchor_left: tuple[int, int]
    anchor_right: tuple[int, int]
    center: tuple[int, int]


def _anchors(x, y, w, h) -> dict:
    cx = x + w // 2
    cy = y + h // 2
    return {
        "anchor_top": (cx, y),
        "anchor_bottom": (cx, y + h),
        "anchor_left": (x, cy),
        "anchor_right": (x + w, cy),
        "center": (cx, cy),
    }


def node(slide, x, y, *,
         label: str, sublabel: str | None = None,
         w: Emu = NODE_W, h: Emu = NODE_H,
         fill: RGBColor = PANEL, border: RGBColor = PANEL_EDGE,
         accent: RGBColor = CYAN,
         label_size: int = 12, sublabel_size: int = 9,
         corner: float = 0.18) -> Node:
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try:
        card.adjustments[0] = corner
    except Exception:
        pass
    card.fill.solid()
    card.fill.fore_color.rgb = fill
    card.line.color.rgb = border
    card.line.width = Pt(0.75)
    card.shadow.inherit = False

    strip = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, x, y, w, Emu(38100)
    )
    strip.line.fill.background()
    fill_gradient(strip, [(0, accent), (1, INDIGO)], angle=0)

    if sublabel:
        text_y = y + Inches(0.12)
        add_text(slide, x + Inches(0.1), text_y, w - Inches(0.2), Inches(0.4),
                 label, size=label_size, color=TEXT, font=FONT_HEAD,
                 bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, x + Inches(0.1), text_y + Inches(0.36),
                 w - Inches(0.2), Inches(0.34),
                 sublabel, size=sublabel_size, color=TEXT_MUTED,
                 font=FONT_MONO, align=PP_ALIGN.CENTER)
    else:
        add_text(slide, x + Inches(0.1), y + Inches(0.18),
                 w - Inches(0.2), h - Inches(0.2),
                 label, size=label_size, color=TEXT, font=FONT_HEAD,
                 bold=True, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)

    return Node(x, y, w, h, card, **_anchors(x, y, w, h))


def arrow(slide,
          frm: Node | tuple[int, int],
          to: Node | tuple[int, int], *,
          side_from: str = "right", side_to: str = "left",
          color: RGBColor = CYAN,
          dashed: bool = False, width: float = 1.75,
          head: bool = True,
          label: str | None = None,
          label_above: bool = True):
    if isinstance(frm, Node):
        fx, fy = getattr(frm, f"anchor_{side_from}")
    else:
        fx, fy = frm
    if isinstance(to, Node):
        tx, ty = getattr(to, f"anchor_{side_to}")
    else:
        tx, ty = to

    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, fx, fy, tx, ty)
    conn.line.color.rgb = color
    conn.line.width = Pt(width)

    ln = conn.line._get_or_add_ln()
    if dashed:
        existing = ln.find(_qn("a:prstDash"))
        if existing is not None:
            ln.remove(existing)
        etree.SubElement(ln, _qn("a:prstDash"), {"val": "dash"})

    if head:
        for tag in ("a:headEnd", "a:tailEnd"):
            el = ln.find(_qn(tag))
            if el is not None:
                ln.remove(el)
        etree.SubElement(ln, _qn("a:tailEnd"),
                         {"type": "triangle", "w": "med", "len": "med"})

    if label:
        mx = (fx + tx) // 2
        my = (fy + ty) // 2
        offset = Inches(-0.24) if label_above else Inches(0.06)
        add_text(slide, mx - Inches(0.9), my + offset,
                 Inches(1.8), Inches(0.25),
                 label, size=9, color=TEXT_DIM, font=FONT_MONO,
                 align=PP_ALIGN.CENTER)
    return conn


def lane(slide, x, y, w, h, *,
         title: str,
         color: RGBColor = INDIGO,
         alpha_fill: RGBColor | None = None):
    """Titled container with a coloured header pill."""
    bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try:
        bg.adjustments[0] = 0.04
    except Exception:
        pass
    bg.fill.solid()
    bg.fill.fore_color.rgb = alpha_fill or BG_MID
    bg.line.color.rgb = color
    bg.line.width = Pt(1.0)

    pill_w = Inches(2.6)
    pill = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.25), y - Inches(0.18),
        pill_w, Inches(0.36),
    )
    pill.adjustments[0] = 0.5
    pill.fill.solid()
    pill.fill.fore_color.rgb = color
    pill.line.fill.background()
    tf = pill.text_frame
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.name = FONT_HEAD
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = BG_DEEP
    return bg


def layer_band(slide, x, y, w, h, *,
               title: str, items: list[str],
               color: RGBColor = INDIGO):
    bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try:
        bg.adjustments[0] = 0.12
    except Exception:
        pass
    bg.fill.solid()
    bg.fill.fore_color.rgb = PANEL
    bg.line.color.rgb = color
    bg.line.width = Pt(1.0)

    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.12), h)
    stripe.line.fill.background()
    fill_gradient(stripe, [(0, color), (1, CYAN)], angle=5400000)

    inner_x = x + Inches(0.35)
    add_text(slide, inner_x, y + Inches(0.10), Inches(3.5), Inches(0.45),
             title, size=15, color=TEXT, font=FONT_HEAD, bold=True)

    items_x = x + Inches(3.95)
    add_text(slide, items_x, y + Inches(0.10), w - Inches(4.15),
             h - Inches(0.2),
             "    ·    ".join(items),
             size=11, color=TEXT_DIM, font=FONT_BODY,
             line_spacing=1.25, anchor=MSO_ANCHOR.MIDDLE)
    return bg
